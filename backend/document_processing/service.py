from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path

from core.config import settings
from .ocr import get_ocr_provider
from storage_backends import get_storage_backend


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg"}


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    ocr_confidence: float | None = None


@dataclass
class ExtractionResult:
    pages: list[ExtractedPage]
    warning: str = ""
    ocr_provider: str = "none"

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)


def store_upload(case_id: str, filename: str, content: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("暂不支持该文件类型。仅允许 PDF、DOCX、TXT、MD、PNG、JPG。")
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise ValueError(f"文件超过 {settings.max_upload_mb}MB 限制。")
    validate_file_signature(filename, content)
    stored = get_storage_backend().put(case_id, filename, content)
    return stored.uri, stored.sha256


def validate_file_signature(filename: str, content: bytes) -> None:
    suffix = Path(filename).suffix.lower()
    if not content:
        raise ValueError("空文件无法处理。")
    signatures = {
        ".pdf": (b"%PDF",), ".docx": (b"PK\x03\x04",),
        ".png": (b"\x89PNG\r\n\x1a\n",), ".jpg": (b"\xff\xd8\xff",), ".jpeg": (b"\xff\xd8\xff",),
    }
    expected = signatures.get(suffix)
    if expected and not any(content.startswith(value) for value in expected):
        raise ValueError("文件扩展名与文件签名不匹配。")
    if suffix in {".txt", ".md"} and b"\x00" in content[:4096]:
        raise ValueError("文本文件包含二进制内容。")


def extract_pages(filename: str, content: bytes) -> ExtractionResult:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return ExtractionResult([ExtractedPage(1, content.decode("utf-8", errors="replace"))])
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            pages = [ExtractedPage(index, page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
            if any(not page.text.strip() for page in pages):
                ocr = get_ocr_provider()
                if ocr.name != "disabled":
                    try:
                        import pypdfium2 as pdfium
                        from PIL import Image

                        pdf = pdfium.PdfDocument(content)
                        for index, page in enumerate(pages):
                            if page.text.strip():
                                continue
                            bitmap = pdf[index].render(scale=2).to_pil()
                            image_buffer = io.BytesIO()
                            bitmap.save(image_buffer, format="PNG")
                            result = ocr.extract(image_buffer.getvalue())
                            page.text = result.text
                            page.ocr_confidence = result.confidence
                        remaining = sum(not page.text.strip() for page in pages)
                        warning = f"仍有 {remaining} 页未提取到文本，需人工核查。" if remaining else "扫描页已使用本地 OCR，需核对识别结果。"
                        return ExtractionResult(pages, warning, ocr.name)
                    except Exception as exc:
                        return ExtractionResult(pages, f"扫描页 OCR 失败，需人工处理：{exc}", ocr.name)
            text = "\n\n".join(page.text for page in pages)
            warning = "" if text.strip() else "PDF 未提取到文本；OCR 未启用或不可用。"
            return ExtractionResult(pages, warning)
        except Exception as exc:  # parser failures are surfaced, never converted into facts
            return ExtractionResult([], f"PDF 解析失败：{exc}")
    if suffix == ".docx":
        try:
            from docx import Document as WordDocument

            document = WordDocument(io.BytesIO(content))
            return ExtractionResult([ExtractedPage(1, "\n".join(p.text for p in document.paragraphs))])
        except Exception as exc:
            return ExtractionResult([], f"Word 解析失败：{exc}")
    ocr = get_ocr_provider().extract(content)
    return ExtractionResult([ExtractedPage(1, ocr.text, ocr.confidence)], ocr.warning, ocr.provider)


def extract_text(filename: str, content: bytes) -> tuple[str, int, str]:
    """Backwards-compatible wrapper used by integrations built against MVP v1."""
    result = extract_pages(filename, content)
    return result.text, len(result.pages), result.warning


def classify_document(filename: str, text: str, case_type: str) -> str:
    haystack = f"{filename} {text[:1000]}".lower()
    rules = [
        (("事故认定", "交通事故认定"), "道路交通事故认定书"),
        (("住院", "入院记录"), "住院病历"),
        (("门诊", "诊断证明"), "门诊病历"),
        (("发票", "票据"), "医疗费用票据" if case_type == "traffic_injury" else "付款凭证"),
        (("费用清单",), "医疗费用清单"),
        (("收入证明", "误工"), "收入及误工证明"),
        (("保险", "保单"), "车辆及保险材料"),
        (("鉴定",), "伤残鉴定材料"),
        (("合同", "协议"), "合同及协议"),
        (("转账", "付款"), "付款凭证"),
        (("聊天", "沟通"), "沟通记录"),
    ]
    for terms, category in rules:
        if any(term in haystack for term in terms):
            return category
    return "其他证据"


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[tuple[int, int, str]]:
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        chunks.append((start, end, text[start:end]))
        if end == len(text):
            break
        start = end - overlap
    return chunks
