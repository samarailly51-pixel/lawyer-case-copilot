from __future__ import annotations

import json
import hashlib
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import desc, select
from sqlalchemy.inspection import inspect
from sqlalchemy.orm import Session

from core.database import get_db
from core.auth import Actor, get_current_actor, require_case_access, require_role
from core.malware import scan_upload
from document_processing.quality import assess_text_quality
from document_processing.service import chunk_text, classify_document, extract_pages, store_upload
from models.entities import (
    AgentRun, AuditLog, Case, CaseParty, CaseReport, CaseTask, CaseWorkspaceLink, CompensationItem, DisabilityAppraisal, Document,
    DocumentChunk, DocumentPage, DocumentQualityAssessment, EvidenceItem, ExtractedFact, HumanReview, InjuryRecord, InsuranceInfo,
    KnowledgeSource, KnowledgeWorkspaceLink, LegalCitation, LegalIssue, MedicalExpense, MissingMaterial, NodeRun, RiskItem, SourceReference,
    TimelineEvent, TrafficAccidentInfo, TrafficRiskItem, TreatmentRecord, PartyRelationship,
)
from schemas.api import (
    BatchReviewCreate, CaseCreate, CaseOut, CaseUpdate, CompensationScenarioRequest,
    DocumentCategoryUpdate, DocumentOut, ReviewCreate, RunCreate, RunResume, TaskUpdate,
)
from schemas.knowledge import KnowledgeSearch, KnowledgeSourceCreate
from services.knowledge import create_knowledge_source, search_knowledge
from services.readiness import production_readiness
from services.rules import validate_rule_file
from evaluation.metrics import evaluate_case
from workflows import WorkflowOrchestrator
from services.workflow_tasks import execute_workflow_run
from core.config import settings
import io
from urllib.parse import quote as url_quote, urlparse
from difflib import SequenceMatcher
import re

from storage_backends import get_storage_backend


router = APIRouter(prefix="/api", dependencies=[Depends(get_current_actor)])


def _value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def row_dict(row: Any, db: Session | None = None, with_sources: bool = False) -> dict[str, Any]:
    result = {column.key: _value(getattr(row, column.key)) for column in inspect(row).mapper.column_attrs}
    if with_sources and db:
        target_names = {
            "ExtractedFact": "fact", "TimelineEvent": "timeline", "EvidenceItem": "evidence",
            "LegalIssue": "legal_issue", "TrafficAccidentInfo": "traffic_accident",
            "InjuryRecord": "injury", "TreatmentRecord": "treatment", "MedicalExpense": "medical_expense",
            "InsuranceInfo": "insurance", "DisabilityAppraisal": "disability_appraisal",
            "TrafficRiskItem": "traffic_risk",
        }
        target_type = target_names.get(type(row).__name__)
        sources = []
        if target_type:
            refs = db.scalars(select(SourceReference).where(SourceReference.target_type == target_type, SourceReference.target_id == row.id))
            for ref in refs:
                doc = db.get(Document, ref.document_id)
                sources.append({"document_id": ref.document_id, "filename": doc.filename if doc else "", "page_number": ref.page_number, "quote": ref.quote})
        result["sources"] = sources
    return result


def current_rows(db: Session, model: Any, case_id: str) -> list[Any]:
    query = select(model).where(model.case_id == case_id)
    if hasattr(model, "is_current"):
        query = query.where(model.is_current.is_(True))
    return list(db.scalars(query.order_by(model.created_at)))


def _storage_key(uri: str) -> str:
    parsed = urlparse(uri)
    if parsed.scheme == "local":
        return f"{parsed.netloc}{parsed.path}".lstrip("/")
    if parsed.scheme == "s3":
        return parsed.path.lstrip("/")
    raise ValueError("材料没有可读取的存储地址")


def _document_page(db: Session, document: Document, page_number: int) -> tuple[str, list[dict[str, Any]]]:
    page = db.scalar(select(DocumentPage).where(
        DocumentPage.document_id == document.id,
        DocumentPage.page_number == page_number,
    ))
    chunks = list(db.scalars(select(DocumentChunk).where(
        DocumentChunk.document_id == document.id,
        DocumentChunk.page_number == page_number,
    ).order_by(DocumentChunk.chunk_index)))
    if page:
        return page.text, [row_dict(chunk) for chunk in chunks]
    if chunks:
        text = "\n".join(chunk.text for chunk in chunks)
        return text, [row_dict(chunk) for chunk in chunks]
    if page_number == 1:
        return document.extracted_text, []
    return "", []


def _locate_quote(text: str, quote: str) -> tuple[str, int, int, bool]:
    if not quote:
        return "", -1, -1, False
    position = text.find(quote)
    if position >= 0:
        return quote, position, position + len(quote), True
    candidates = [item.strip() for item in re.split(r"(?<=[。！？；])|\n", text) if len(item.strip()) >= 6]
    if not candidates:
        return quote, -1, -1, False
    best = max(candidates, key=lambda item: SequenceMatcher(None, item, quote).ratio())
    score = SequenceMatcher(None, best, quote).ratio()
    if score < 0.28:
        return quote, -1, -1, False
    position = text.find(best)
    return best, position, position + len(best), False


def _report_material_references(db: Session, case_id: str) -> list[dict[str, Any]]:
    refs = list(db.scalars(select(SourceReference).where(SourceReference.case_id == case_id).order_by(
        SourceReference.document_id, SourceReference.page_number, SourceReference.created_at
    )))
    result: list[dict[str, Any]] = []
    seen: set[tuple[str, int | None, str]] = set()
    for ref in refs:
        key = (ref.document_id, ref.page_number, ref.quote.strip())
        if key in seen or not ref.quote.strip():
            continue
        seen.add(key)
        document = db.get(Document, ref.document_id)
        result.append({
            "index": len(result) + 1,
            "document_id": ref.document_id,
            "filename": document.filename if document else "未知材料",
            "page_number": ref.page_number or 1,
            "quote": ref.quote.strip(),
            "target_type": ref.target_type,
            "target_id": ref.target_id,
        })
    return result


@router.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "product": "Lawyer Case Copilot"}


@router.get("/cases", response_model=list[CaseOut])
def list_cases(actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    return list(db.scalars(
        select(Case).join(CaseWorkspaceLink, CaseWorkspaceLink.case_id == Case.id)
        .where(CaseWorkspaceLink.workspace_id == actor.workspace_id).order_by(desc(Case.updated_at))
    ))


@router.get("/portfolio-metrics")
def portfolio_metrics(actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    """Return evidence-backed portfolio metrics for the synthetic demo workspace."""
    cases = list(db.scalars(
        select(Case).join(CaseWorkspaceLink, CaseWorkspaceLink.case_id == Case.id)
        .where(CaseWorkspaceLink.workspace_id == actor.workspace_id, Case.is_demo.is_(True))
        .order_by(Case.created_at)
    ))
    reports = [evaluate_case(db, item.id) for item in cases]
    case_ids = [item.id for item in cases]
    if not case_ids:
        return {
            "dataset_label": "完全虚构回归集",
            "case_count": 0,
            "document_count": 0,
            "fact_count": 0,
            "fact_source_coverage": 0.0,
            "compensation_item_count": 0,
            "workflow_node_count": 0,
            "mandatory_risk_count": 0,
            "unsupported_model_outputs_rejected": 0,
            "disclaimer": "仅衡量系统输出完整性和可追溯性，不评价法律结论。",
        }
    compensation_items = list(db.scalars(select(CompensationItem).where(
        CompensationItem.case_id.in_(case_ids), CompensationItem.is_current.is_(True)
    )))
    mandatory_risks = list(db.scalars(select(TrafficRiskItem).where(
        TrafficRiskItem.case_id.in_(case_ids),
        TrafficRiskItem.is_current.is_(True),
        TrafficRiskItem.mandatory_human_review.is_(True),
    )))
    node_names = set(db.scalars(select(NodeRun.node_name).where(NodeRun.case_id.in_(case_ids))))
    return {
        "dataset_label": "完全虚构回归集",
        "case_count": len(cases),
        "document_count": sum(report.counts["documents"] for report in reports),
        "fact_count": sum(report.counts["facts"] for report in reports),
        "fact_source_coverage": round(sum(report.fact_source_coverage for report in reports) / len(reports), 4),
        "compensation_item_count": len(compensation_items),
        "workflow_node_count": len(node_names),
        "mandatory_risk_count": len(mandatory_risks),
        "unsupported_model_outputs_rejected": sum(report.unsupported_model_outputs_rejected for report in reports),
        "disclaimer": "仅衡量完全虚构样例的系统输出完整性和可追溯性，不评价法律结论或案件结果。",
    }


@router.get("/evaluation/latest")
def latest_evaluation(actor: Actor = Depends(get_current_actor)):
    """Expose the checked-in synthetic regression evidence without running jobs in a request."""
    result_path = Path(__file__).resolve().parents[2] / "evals" / "latest-results.json"
    if not result_path.exists():
        return {
            "generated_at": None,
            "summary": {
                "dataset_label": "完全虚构的 Demo 与非 Demo 结构化回归集",
                "total_scenarios": 0,
                "passed_scenarios": 0,
                "pass_rate": 0,
                "average_fact_source_coverage": 0,
                "scope_note": "尚未生成评测快照，请运行 python -m evaluation.run_suite。",
            },
            "results": [],
            "limitations": [
                "该评测只验证结构完整性、材料忠实性、可追溯性和人工复核边界。",
                "它不代表真实案件准确率、法律正确率、医疗判断或业务效果。",
            ],
        }
    try:
        payload = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(503, f"评测快照不可用：{exc}") from exc
    payload["limitations"] = [
        "全部评测材料均为完全虚构的合成样例。",
        "100% 仅表示当前合成回归断言通过，不代表真实案件准确率。",
        "真实脱敏数据集、律师标注阈值和业务效果仍须由项目所有者补充。",
    ]
    return payload


@router.post("/cases", response_model=CaseOut, status_code=201)
def create_case(payload: CaseCreate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_role(actor, "lawyer")
    case = Case(**payload.model_dump())
    db.add(case); db.flush()
    db.add(CaseWorkspaceLink(case_id=case.id, workspace_id=actor.workspace_id))
    db.commit(); db.refresh(case)
    if payload.client_name:
        db.add(CaseParty(case_id=case.id, name=payload.client_name, role="委托人", created_by_type="user", review_status="accepted"))
    if payload.opposing_party:
        db.add(CaseParty(case_id=case.id, name=payload.opposing_party, role="对方当事人", created_by_type="user", review_status="accepted"))
    db.commit()
    return case


@router.get("/cases/{case_id}", response_model=CaseOut)
def get_case(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    case = db.get(Case, case_id)
    if not case:
        raise HTTPException(404, "案件不存在")
    return case


@router.patch("/cases/{case_id}", response_model=CaseOut)
def update_case(case_id: str, payload: CaseUpdate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "lawyer")
    case = db.get(Case, case_id)
    if not case:
        raise HTTPException(404, "案件不存在")
    for key, value in payload.model_dump(exclude_none=True).items():
        setattr(case, key, value)
    db.add(AuditLog(case_id=case.id, actor=actor.display_name, action="update_case", target_type="case", target_id=case.id, details=payload.model_dump(exclude_none=True)))
    db.commit(); db.refresh(case)
    return case


@router.get("/cases/{case_id}/documents", response_model=list[DocumentOut])
def list_documents(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    return list(db.scalars(select(Document).where(Document.case_id == case_id).order_by(Document.created_at)))


@router.post("/cases/{case_id}/documents", response_model=DocumentOut, status_code=201)
async def upload_document(case_id: str, file: UploadFile = File(...), actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "assistant")
    case = db.get(Case, case_id)
    if not case:
        raise HTTPException(404, "案件不存在")
    content = await file.read()
    try:
        malware_scan = scan_upload(content)
        if not malware_scan.clean:
            raise ValueError(f"文件未通过恶意内容扫描：{malware_scan.signature or '检测到风险'}")
        target, digest = store_upload(case_id, file.filename or "upload.txt", content)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    extraction = extract_pages(file.filename or "upload.txt", content)
    text, pages, warning = extraction.text, len(extraction.pages), extraction.warning
    category = classify_document(file.filename or "upload.txt", text, case.case_type)
    doc = Document(case_id=case_id, filename=file.filename or "upload", mime_type=file.content_type or "application/octet-stream",
                   file_path=str(target), sha256=digest, category=category, classification_source="ai",
                   parse_status="parsed" if text else "needs_review", parse_warning=warning,
                   extracted_text=text, page_count=pages or 1)
    db.add(doc); db.flush()
    chunk_index = 0
    for page in extraction.pages:
        db.add(DocumentPage(
            document_id=doc.id, page_number=page.page_number, text=page.text,
            source_mode=page.source_mode, ocr_confidence=page.ocr_confidence,
            ocr_regions=page.ocr_regions or [], image_width=page.image_width,
            image_height=page.image_height,
        ))
        for start, end, value in chunk_text(page.text):
            db.add(DocumentChunk(document_id=doc.id, page_number=page.page_number, chunk_index=chunk_index, text=value,
                                 start_offset=start, end_offset=end, ocr_confidence=page.ocr_confidence))
            chunk_index += 1
    ocr_scores = [page.ocr_confidence for page in extraction.pages if page.ocr_confidence is not None]
    average_ocr = sum(ocr_scores) / len(ocr_scores) if ocr_scores else None
    quality = assess_text_quality(text, warning, average_ocr)
    db.add(DocumentQualityAssessment(
        document_id=doc.id, case_id=case_id, text_quality_score=quality.score,
        injection_risk=quality.injection_risk, security_flags=list(dict.fromkeys([*quality.flags, *malware_scan.flags])),
        warnings=list(dict.fromkeys([*quality.warnings, *malware_scan.warnings])),
        requires_human_review=quality.requires_human_review or bool(malware_scan.flags),
        ocr_provider=extraction.ocr_provider,
    ))
    db.add(AuditLog(
        case_id=case_id,
        actor=actor.display_name,
        action="upload_document",
        target_type="document",
        target_id=doc.id,
        details={
            "filename": doc.filename,
            "sha256": digest,
            "malware_scan_provider": malware_scan.provider,
            "malware_scan_clean": malware_scan.clean,
        },
    ))
    db.commit(); db.refresh(doc)
    return doc


@router.patch("/documents/{document_id}/classification", response_model=DocumentOut)
def update_document_category(document_id: str, payload: DocumentCategoryUpdate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    doc = db.get(Document, document_id)
    if not doc:
        raise HTTPException(404, "材料不存在")
    require_case_access(db, actor, doc.case_id, "assistant")
    old = doc.category
    doc.category = payload.category
    doc.classification_source = "user"
    db.add(AuditLog(case_id=doc.case_id, actor=actor.display_name, action="reclassify_document",
                    target_type="document", target_id=doc.id, details={"from": old, "to": payload.category}))
    db.commit(); db.refresh(doc)
    return doc


@router.get("/documents/{document_id}/preview")
def preview_document(document_id: str, page: int = Query(default=1, ge=1),
                     highlight: str = Query(default="", max_length=1000),
                     actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    document = db.get(Document, document_id)
    if not document:
        raise HTTPException(404, "材料不存在")
    require_case_access(db, actor, document.case_id)
    if page > max(document.page_count, 1):
        raise HTTPException(404, "材料页码不存在")
    text, chunks = _document_page(db, document, page)
    page_record = db.scalar(select(DocumentPage).where(
        DocumentPage.document_id == document.id,
        DocumentPage.page_number == page,
    ))
    matched, position, end_position, exact = _locate_quote(text, highlight)
    return {
        "id": document.id, "filename": document.filename, "mime_type": document.mime_type,
        "page_number": page, "page_count": max(document.page_count, 1), "text": text,
        "requested_highlight": highlight, "highlight": matched, "highlight_exact": exact,
        "highlight_start": position, "highlight_end": end_position,
        "chunks": chunks, "has_original": bool(document.file_path),
        "parse_warning": document.parse_warning,
        "page_quality": {
            "source_mode": page_record.source_mode if page_record else "legacy",
            "ocr_confidence": page_record.ocr_confidence if page_record else next(
                (item.get("ocr_confidence") for item in chunks if item.get("ocr_confidence") is not None), None
            ),
            "ocr_regions": page_record.ocr_regions if page_record else [],
            "image_width": page_record.image_width if page_record else None,
            "image_height": page_record.image_height if page_record else None,
        },
    }


@router.get("/documents/{document_id}/content")
def document_content(document_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    document = db.get(Document, document_id)
    if not document:
        raise HTTPException(404, "材料不存在")
    require_case_access(db, actor, document.case_id)
    if not document.file_path:
        raise HTTPException(404, "该演示材料没有原始二进制文件，请使用文本预览")
    try:
        content = get_storage_backend().get(_storage_key(document.file_path))
    except (ValueError, FileNotFoundError) as exc:
        raise HTTPException(404, "原始材料不可用") from exc
    return Response(content, media_type=document.mime_type, headers={
        "Content-Disposition": f"inline; filename*=UTF-8''{url_quote(document.filename)}",
        "X-Content-Type-Options": "nosniff",
        "Content-Security-Policy": "sandbox",
    })


@router.post("/cases/{case_id}/runs", status_code=201)
def create_run(case_id: str, payload: RunCreate, background_tasks: BackgroundTasks,
               actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "assistant")
    try:
        run = WorkflowOrchestrator(db).create_run(case_id, payload.trigger_type)
        if settings.workflow_execution_mode == "background":
            background_tasks.add_task(execute_workflow_run, run.id)
        elif settings.workflow_execution_mode == "inline":
            run = WorkflowOrchestrator(db).execute(run.id)
        return row_dict(run)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc


@router.get("/cases/{case_id}/runs")
def list_runs(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    runs = list(db.scalars(select(AgentRun).where(AgentRun.case_id == case_id).order_by(desc(AgentRun.created_at))))
    return [row_dict(run) for run in runs]


@router.get("/runs/{run_id}")
def get_run(run_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    run = db.get(AgentRun, run_id)
    if not run:
        raise HTTPException(404, "运行记录不存在")
    require_case_access(db, actor, run.case_id)
    nodes = list(db.scalars(select(NodeRun).where(NodeRun.run_id == run_id).order_by(NodeRun.sequence, NodeRun.created_at)))
    serialized_nodes = []
    attempts: dict[str, int] = {}
    total_duration_ms = 0
    for node in nodes:
        item = row_dict(node)
        duration_ms = None
        if node.started_at and node.completed_at:
            duration_ms = max(0, round((node.completed_at - node.started_at).total_seconds() * 1000))
            total_duration_ms += duration_ms
        item["duration_ms"] = duration_ms
        attempts[node.node_name] = attempts.get(node.node_name, 0) + 1
        item["attempt"] = attempts[node.node_name]
        serialized_nodes.append(item)
    run_duration_ms = None
    if run.started_at and run.completed_at:
        run_duration_ms = max(0, round((run.completed_at - run.started_at).total_seconds() * 1000))
    return {
        **row_dict(run),
        "duration_ms": run_duration_ms,
        "node_duration_ms": total_duration_ms,
        "node_counts": {
            "completed": sum(item["status"] == "completed" for item in serialized_nodes),
            "failed": sum(item["status"] == "failed" for item in serialized_nodes),
            "skipped": sum(item["status"] == "skipped" for item in serialized_nodes),
            "rerun": sum(bool(item["supersedes_node_run_id"]) for item in serialized_nodes),
        },
        "nodes": serialized_nodes,
    }


@router.post("/runs/{run_id}/nodes/{node_name}/rerun")
def rerun_node(run_id: str, node_name: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    run = db.get(AgentRun, run_id)
    if not run:
        raise HTTPException(404, "运行记录不存在")
    require_case_access(db, actor, run.case_id, "assistant")
    try:
        return row_dict(WorkflowOrchestrator(db).rerun_node(run_id, node_name))
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@router.post("/runs/{run_id}/resume")
def resume_run(run_id: str, payload: RunResume, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    run = db.get(AgentRun, run_id)
    if not run:
        raise HTTPException(404, "运行记录不存在")
    require_case_access(db, actor, run.case_id, "assistant")
    try:
        resumed = WorkflowOrchestrator(db).rerun_from_node(run_id, payload.from_node)
        db.add(AuditLog(
            case_id=run.case_id, actor=actor.display_name, action="resume_workflow",
            target_type="agent_run", target_id=run.id,
            details={"from_node": payload.from_node or "failed_node", "status": resumed.status},
        ))
        db.commit()
        return row_dict(resumed)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@router.get("/cases/{case_id}/workspace")
def case_workspace(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    case = db.get(Case, case_id)
    if not case:
        raise HTTPException(404, "案件不存在")
    model_map = {
        "facts": ExtractedFact, "timeline": TimelineEvent, "evidence": EvidenceItem,
        "issues": LegalIssue, "missing_materials": MissingMaterial, "risks": RiskItem,
        "tasks": CaseTask, "compensation_items": CompensationItem, "traffic_risks": TrafficRiskItem,
        "injuries": InjuryRecord, "treatments": TreatmentRecord, "medical_expenses": MedicalExpense,
        "insurance": InsuranceInfo, "traffic_accident": TrafficAccidentInfo, "reports": CaseReport,
        "disability_appraisals": DisabilityAppraisal,
        "legal_citations": LegalCitation, "document_quality": DocumentQualityAssessment,
        "relationships": PartyRelationship,
    }
    return {
        "case": CaseOut.model_validate(case).model_dump(mode="json"),
        "documents": [row_dict(item) for item in current_rows(db, Document, case_id)],
        **{key: [row_dict(item, db, True) for item in current_rows(db, model, case_id)] for key, model in model_map.items()},
        "parties": [row_dict(item) for item in current_rows(db, CaseParty, case_id)],
    }


@router.post("/knowledge/sources", status_code=201)
def add_knowledge_source(payload: KnowledgeSourceCreate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_role(actor, "lawyer")
    item = create_knowledge_source(db, payload)
    link = db.scalar(select(KnowledgeWorkspaceLink).where(
        KnowledgeWorkspaceLink.knowledge_source_id == item.id, KnowledgeWorkspaceLink.workspace_id == actor.workspace_id
    ))
    if not link:
        db.add(KnowledgeWorkspaceLink(knowledge_source_id=item.id, workspace_id=actor.workspace_id)); db.commit()
    return row_dict(item)


@router.get("/knowledge/sources")
def list_knowledge_sources(scope: str | None = None, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    query = select(KnowledgeSource).join(
        KnowledgeWorkspaceLink, KnowledgeWorkspaceLink.knowledge_source_id == KnowledgeSource.id
    ).where(KnowledgeWorkspaceLink.workspace_id == actor.workspace_id).order_by(desc(KnowledgeSource.published_or_updated_at))
    if scope:
        query = query.where(KnowledgeSource.scope == scope)
    return [row_dict(item) for item in db.scalars(query)]


@router.post("/knowledge/search")
def knowledge_search(payload: KnowledgeSearch, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    return [row_dict(item) for item in search_knowledge(
        db, payload.query, payload.scopes, payload.limit, actor.workspace_id,
        payload.verified_only, payload.exclude_historical,
    )]


@router.post("/cases/{case_id}/legal-retrieval")
def retrieve_for_case(case_id: str, payload: KnowledgeSearch, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "assistant")
    if not db.get(Case, case_id):
        raise HTTPException(404, "案件不存在")
    sources = search_knowledge(
        db, payload.query, payload.scopes, payload.limit, actor.workspace_id,
        payload.verified_only, payload.exclude_historical,
    )
    citations = []
    for source in sources:
        item = LegalCitation(
            case_id=case_id, title=source.title, excerpt=source.excerpt, source_name=source.source_name,
            source_url=source.source_url, published_or_updated_at=source.published_or_updated_at,
            jurisdiction=source.jurisdiction, scope=source.applicability_scope,
            stale_risk=source.stale_risk, confidence=1.0, created_by_type="system",
        )
        db.add(item); db.flush(); citations.append(item)
    db.commit()
    return [row_dict(item) for item in citations]


@router.get("/cases/{case_id}/quality")
def case_quality(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    try:
        return evaluate_case(db, case_id).to_dict()
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc


REVIEW_TARGETS = {
    "fact": ExtractedFact, "timeline": TimelineEvent, "evidence": EvidenceItem,
    "legal_issue": LegalIssue, "missing_material": MissingMaterial, "risk": RiskItem,
    "traffic_risk": TrafficRiskItem, "compensation_item": CompensationItem, "report": CaseReport,
}


def _apply_review(db: Session, actor: Actor, payload: ReviewCreate) -> HumanReview:
    model = REVIEW_TARGETS.get(payload.target_type)
    if not model:
        raise HTTPException(400, "不支持的复核对象类型")
    target = db.get(model, payload.target_id)
    if not target or target.case_id != payload.case_id:
        raise HTTPException(404, "复核对象不存在")
    original = row_dict(target)
    target.review_status = payload.action
    allowed = {column.key for column in inspect(target).mapper.column_attrs} - {
        "id", "case_id", "created_at", "updated_at", "agent_run_id", "version", "is_current"
    }
    if payload.action == "modified":
        for key, value in payload.revised_value.items():
            if key in allowed:
                setattr(target, key, value)
        target.created_by_type = "user"
        target.version += 1
    review_data = payload.model_dump(exclude={"reviewer"})
    review = HumanReview(**review_data, reviewer=actor.display_name, original_value=original)
    db.add(review)
    db.add(AuditLog(case_id=payload.case_id, actor=actor.display_name, action=f"review_{payload.action}",
                    target_type=payload.target_type, target_id=payload.target_id,
                    details={"comment": payload.comment, "revised": payload.revised_value, "version": target.version}))
    return review


@router.post("/reviews", status_code=201)
def create_review(payload: ReviewCreate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, payload.case_id, "lawyer")
    review = _apply_review(db, actor, payload)
    db.commit(); db.refresh(review)
    return row_dict(review)


@router.post("/reviews/batch", status_code=201)
def create_batch_reviews(payload: BatchReviewCreate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, payload.case_id, "lawyer")
    reviews: list[HumanReview] = []
    for item in payload.items:
        if not item.get("target_type") or not item.get("target_id"):
            raise HTTPException(400, "批量复核项缺少 target_type 或 target_id")
        reviews.append(_apply_review(db, actor, ReviewCreate(
            case_id=payload.case_id, target_type=item["target_type"], target_id=item["target_id"],
            action=payload.action, comment=payload.comment, error_type=payload.error_type,
        )))
    db.commit()
    return {"processed": len(reviews), "review_ids": [item.id for item in reviews]}


@router.get("/reviews/{target_type}/{target_id}/history")
def review_history(target_type: str, target_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    model = REVIEW_TARGETS.get(target_type)
    target = db.get(model, target_id) if model else None
    if not target:
        raise HTTPException(404, "复核对象不存在")
    require_case_access(db, actor, target.case_id)
    reviews = list(db.scalars(select(HumanReview).where(
        HumanReview.case_id == target.case_id,
        HumanReview.target_type == target_type,
        HumanReview.target_id == target_id,
    ).order_by(HumanReview.created_at)))
    return {"current": row_dict(target, db, True), "history": [row_dict(item) for item in reviews]}


@router.patch("/tasks/{task_id}")
def update_task(task_id: str, payload: TaskUpdate, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    task = db.get(CaseTask, task_id)
    if not task:
        raise HTTPException(404, "任务不存在")
    require_case_access(db, actor, task.case_id, "assistant")
    changes = payload.model_dump(exclude_none=True)
    for key, value in changes.items():
        setattr(task, key, value)
    task.review_status = "modified"
    task.created_by_type = "user"
    db.add(AuditLog(case_id=task.case_id, actor=actor.display_name, action="update_task",
                    target_type="task", target_id=task.id, details=changes))
    db.commit(); db.refresh(task)
    return row_dict(task)


@router.get("/cases/{case_id}/audit-logs")
def case_audit_logs(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "lawyer")
    return [row_dict(item) for item in db.scalars(select(AuditLog).where(
        AuditLog.case_id == case_id
    ).order_by(desc(AuditLog.created_at)).limit(200))]


@router.get("/cases/{case_id}/reviews")
def list_reviews(case_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    return [row_dict(item) for item in db.scalars(select(HumanReview).where(HumanReview.case_id == case_id).order_by(desc(HumanReview.created_at)))]


@router.get("/cases/{case_id}/reports/{report_id}/citations")
def report_citations(case_id: str, report_id: str, actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    report = db.get(CaseReport, report_id)
    if not report or report.case_id != case_id or not report.is_current:
        raise HTTPException(404, "报告不存在")
    return _report_material_references(db, case_id)


@router.get("/traffic-injury/rules")
def traffic_rule_catalog(actor: Actor = Depends(get_current_actor)):
    root = Path(__file__).resolve().parents[1] / "rules" / "traffic_injury"
    catalog = []
    for filename in ("evidence_completeness.yaml", "personal_experience_rules.yaml"):
        validation = validate_rule_file(root / filename)
        payload = validation.payload
        file_sha256 = hashlib.sha256((root / filename).read_bytes()).hexdigest()
        catalog.append({
            "filename": filename, "version": payload.get("version", 1),
            "file_sha256": file_sha256,
            "description": payload.get("description", "系统证据完整性规则"),
            "rule_count": len(payload.get("rules") or []), "rules": payload.get("rules") or [],
            "enabled_rule_count": validation.enabled_rule_count,
            "valid": validation.valid,
            "validation_errors": validation.errors,
            "validation_warnings": validation.warnings,
        })
    return catalog


@router.get("/system/readiness")
def system_readiness(actor: Actor = Depends(get_current_actor)):
    require_role(actor, "admin")
    return production_readiness()


@router.post("/cases/{case_id}/compensation-scenario")
def compensation_scenario(case_id: str, payload: CompensationScenarioRequest,
                          actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id, "assistant")
    case = db.get(Case, case_id)
    if not case or case.case_type != "traffic_injury":
        raise HTTPException(400, "仅交通事故人伤案件可使用情景测算")
    p = payload.parameters
    calculations = [
        ("医疗费", p.get("medical_expense")),
        ("误工费", p.get("lost_work_days", 0) * p.get("lost_work_daily_rate", 0)),
        ("护理费", p.get("nursing_days", 0) * p.get("nursing_daily_rate", 0)),
        ("营养费", p.get("nutrition_days", 0) * p.get("nutrition_daily_rate", 0)),
        ("住院伙食补助费", p.get("hospital_days", 0) * p.get("hospital_daily_rate", 0)),
        ("交通费", p.get("transportation_expense")),
        ("鉴定费", p.get("appraisal_expense")),
        ("辅助器具费", p.get("assistive_device_expense")),
    ]
    items = [{"name": name, "amount": round(float(amount), 2), "confirmed": False}
             for name, amount in calculations if amount is not None and float(amount) != 0]
    return {
        "items": items, "total": round(sum(item["amount"] for item in items), 2),
        "parameters": p, "status": "scenario_only",
        "warnings": [
            "本结果仅按律师录入参数进行算术汇总，不代表项目成立、标准适用或金额可获支持。",
            "地区标准、责任比例、保险限额、证据效力及其他调整项尚未自动适用。",
            "所有参数与结果均须由案件负责律师核验。",
        ],
    }


@router.get("/cases/{case_id}/reports/{report_id}/export")
def export_report(case_id: str, report_id: str, format: str = Query(default="docx", pattern="^(docx|md)$"),
                  actor: Actor = Depends(get_current_actor), db: Session = Depends(get_db)):
    require_case_access(db, actor, case_id)
    report = db.get(CaseReport, report_id)
    if not report or report.case_id != case_id or not report.is_current:
        raise HTTPException(404, "报告不存在")
    safe_name = "案件辅助报告"
    references = _report_material_references(db, case_id)
    reference_markdown = "\n".join(
        f"[{item['index']}] {item['filename']}，第 {item['page_number']} 页：{item['quote']}" for item in references
    ) or "当前报告没有可用的材料引用。"
    if format == "md":
        content = f"{report.content}\n\n## 材料引用目录\n\n{reference_markdown}\n\n---\n\n{report.disclaimer}\n".encode("utf-8")
        return Response(content, media_type="text/markdown; charset=utf-8", headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{url_quote(safe_name)}.md"
        })
    from docx import Document as WordDocument
    document = WordDocument()
    document.add_heading(report.title, level=1)
    for line in report.content.splitlines():
        value = line.strip()
        if value.startswith("# "):
            continue
        if value.startswith("## "):
            document.add_heading(value[3:], level=2)
        elif value.startswith("- "):
            document.add_paragraph(value[2:], style="List Bullet")
        elif value.startswith("> "):
            document.add_paragraph(value[2:], style="Intense Quote")
        elif value:
            document.add_paragraph(value)
    document.add_heading("材料引用目录", level=2)
    for item in references:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(f"{item['filename']}，第 {item['page_number']} 页：").bold = True
        paragraph.add_run(item["quote"])
    if not references:
        document.add_paragraph("当前报告没有可用的材料引用。")
    document.add_paragraph(report.disclaimer, style="Intense Quote")
    buffer = io.BytesIO(); document.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{url_quote(safe_name)}.docx"})
